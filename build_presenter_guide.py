from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("/Users/mandinu/Downloads/project/output/presentation/SPHERE_Four_Presenter_Guide.docx")

NAVY = "16324F"
BLUE = "2D6A9F"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F3F5F7"
MID_GRAY = "D9D9D9"
TEXT = RGBColor(35, 42, 48)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, keep=True):
    paragraph.paragraph_format.keep_with_next = keep


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_label_paragraph(doc, label, text, space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(label)
    r.bold = True
    r.font.color.rgb = TEXT
    p.add_run(text)
    return p


def add_bullet(doc, text, level=0, space_after=3):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.05
    return p


def add_number(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.06
    return p


def add_script(doc, paragraphs):
    for text in paragraphs:
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        p.paragraph_format.keep_together = True


def add_role_table(doc):
    rows = [
        ("Elio", "Opening and objectives", "Slides 1–4", "2:30–3:00", "Why the project exists and what the audience will see"),
        ("Dilan", "Experiment and physics", "Slides 5–8", "3:15–3:45", "Hardware, insulation builds, model, method and safety"),
        ("Piyath", "Mission Control and demo", "Slides 9–11", "3:45–4:45", "How prediction and measurement connect in the web tool"),
        ("Nuhansee", "Evidence and closing", "Slides 12–14", "2:45–3:15", "Results, correct interpretation, project offer and questions"),
    ]
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.85, 1.35, 0.85, 0.9, 3.0]
    headers = ["Presenter", "Section", "Slides", "Time", "Main responsibility"]
    for i, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Inches(widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.5)
    set_repeat_table_header(table.rows[0])
    for r_idx, row_data in enumerate(rows, 1):
        cells = table.add_row().cells
        for i, (cell, text) in enumerate(zip(cells, row_data)):
            cell.width = Inches(widths[i])
            set_cell_margins(cell)
            set_cell_shading(cell, PALE_BLUE if r_idx % 2 == 0 else "FFFFFF")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            if i == 0:
                run.bold = True
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_qa_table(doc):
    rows = [
        ("Project purpose, audience and learning objectives", "Elio"),
        ("Thermodynamics, hardware, safety and experimental controls", "Dilan"),
        ("Mission Control, model settings, telemetry and browser demo", "Piyath"),
        ("Results, uncertainty, classroom package and next steps", "Nuhansee"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [5.6, 1.5]
    for i, (cell, text) in enumerate(zip(table.rows[0].cells, ["Question area", "First responder"])):
        cell.width = Inches(widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    for r_idx, row_data in enumerate(rows, 1):
        cells = table.add_row().cells
        for i, (cell, text) in enumerate(zip(cells, row_data)):
            cell.width = Inches(widths[i])
            set_cell_margins(cell)
            set_cell_shading(cell, PALE_GRAY if r_idx % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            if i == 1:
                r.bold = True
    set_table_borders(table)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = TEXT
styles["Normal"].paragraph_format.space_after = Pt(5)

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(28)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(9)
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for style_name, size in (("Heading 1", 17), ("Heading 2", 12.5)):
    s = styles[style_name]
    s.font.name = "Aptos Display"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(9 if style_name == "Heading 1" else 6)
    s.paragraph_format.space_after = Pt(4)

for style_name in ("List Bullet", "List Bullet 2", "List Number"):
    styles[style_name].font.name = "Aptos"
    styles[style_name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles[style_name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles[style_name].font.size = Pt(10.5)

# Cover and overview
p = doc.add_paragraph("SPHERE Four Presenter Guide", style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p = doc.add_paragraph("SPACESUIT THERMAL SHIELDING KIT final presentation")
p.paragraph_format.space_after = Pt(14)
r = p.runs[0]
r.font.size = Pt(15)
r.font.bold = True
r.font.color.rgb = RGBColor(*[int(BLUE[i:i+2], 16) for i in (0, 2, 4)])

add_label_paragraph(doc, "Purpose  ", "This guide gives the four presenters one clear story, balanced speaking roles, natural handoffs and a shared answer to the question: what does SPHERE offer?")
add_label_paragraph(doc, "Target length  ", "Approximately 13 to 15 minutes, followed by questions. Appendix slides remain available only for questions.")

add_heading(doc, "The project in one sentence", 1)
p = doc.add_paragraph("SPHERE is a reusable classroom thermal laboratory that lets students predict, build, measure and evaluate a spacesuit-inspired insulation system using simple hardware and browser-based Mission Control.")
p.paragraph_format.space_after = Pt(8)
p.paragraph_format.line_spacing = 1.1

add_heading(doc, "Presentation flow", 1)
add_number(doc, "1.  Problem and objective: explain why thermal protection matters and what students learn.")
add_number(doc, "2.  Physical experiment: show the four builds, the heat-transfer model, fair testing and safety.")
add_number(doc, "3.  Digital workflow: demonstrate how Mission Control links predicted and measured cooling.")
add_number(doc, "4.  Evidence and offer: interpret the results carefully, then close with the complete classroom package.")

add_heading(doc, "Four speaking roles", 1)
add_role_table(doc)

doc.add_page_break()

add_heading(doc, "Shared message for all presenters", 1)
add_heading(doc, "Objectives", 2)
add_bullet(doc, "Students distinguish conduction, convection and radiation in a practical system.")
add_bullet(doc, "Students use Newton's cooling model and interpret the effective cooling constant k.")
add_bullet(doc, "Students collect a controlled 16-point temperature series and compare it with a prediction.")
add_bullet(doc, "Students explain model disagreement using uncertainty and physical heat-transfer mechanisms.")

add_heading(doc, "What SPHERE offers", 2)
add_bullet(doc, "A complete theory-to-evidence learning loop: predict, build, measure, compare and critique.")
add_bullet(doc, "A reusable four-group experiment with simple, accessible hardware.")
add_bullet(doc, "Mission Control, a browser application for prediction, telemetry entry, graphs and error analysis.")
add_bullet(doc, "Aligned manuals, illustrated SOP, task sheet, revision flashcards and ten-question assessment.")
add_bullet(doc, "A classroom-ready approach that requires no native software installation.")

add_heading(doc, "Three accuracy rules", 2)
add_number(doc, "1.  Call the ice-water setup a classroom analogy for spacecraft thermal control. Do not call it a vacuum simulation.")
add_number(doc, "2.  Describe k as an effective cooling coefficient for the complete tested system. Do not call it the thermal conductivity of a material.")
add_number(doc, "3.  Describe the measurements as a worked example. One run supports discussion but does not validate every material or build.")

add_heading(doc, "Delivery style", 2)
add_bullet(doc, "Explain the purpose of each slide instead of reading every line.")
add_bullet(doc, "Face the audience for the main message; turn to the screen only to point at evidence.")
add_bullet(doc, "End each section with one complete conclusion, pause, then name the next presenter.")
add_bullet(doc, "If time is short, keep the objective, the demo result, the measured ranking and the final offer.")

doc.add_page_break()

# Presenter 1
add_heading(doc, "Presenter 1 Elio", 1)
add_label_paragraph(doc, "Slides  ", "1 to 4")
add_label_paragraph(doc, "Goal  ", "Give the audience a reason to care, define the project accurately and make the learning objectives easy to remember.")
add_label_paragraph(doc, "Target time  ", "2 minutes 30 seconds to 3 minutes")

add_heading(doc, "Speaking script", 2)
add_script(doc, [
    "Good morning. We are presenting the SPACESUIT THERMAL SHIELDING KIT, inspired by the challenge of protecting an astronaut from extreme temperature changes. Our project turns that engineering problem into a measurable classroom experiment.",
    "The central question is simple: how do layered materials change the cooling rate of a protected thermal mass? Warm water represents the protected core, an ice-water bath creates a repeatable cold boundary, and four insulation builds let students compare different heat-transfer strategies.",
    "SPHERE is more than one experiment. It combines the physical kit with Mission Control, student and teacher manuals, an illustrated procedure, a task sheet, revision flashcards and an aligned assessment. Students use the same settings and terminology from preparation through evaluation.",
    "The learning objectives connect directly to student actions. Students identify conduction, convection and radiation, use Newton's cooling model, collect a 16-point temperature history and evaluate the difference between prediction and observation. Four groups can work in parallel in one classroom session.",
])

add_heading(doc, "Points to emphasize", 2)
add_bullet(doc, "The project is designed for learning through measurement, not only for explaining theory.")
add_bullet(doc, "The cold-water test is an analogy that makes heat transfer visible in class.")
add_bullet(doc, "The full package links the experiment, software and assessment.")

add_heading(doc, "Handoff to Dilan", 2)
p = doc.add_paragraph("To show how we turn these objectives into a controlled and safe experiment, Dilan will explain the hardware, insulation configurations and thermodynamic model.")
p.paragraph_format.keep_together = True

doc.add_page_break()

# Presenter 2
add_heading(doc, "Presenter 2 Dilan", 1)
add_label_paragraph(doc, "Slides  ", "5 to 8")
add_label_paragraph(doc, "Goal  ", "Explain how the experiment works, why each insulation build is different and how the team controls safety and fairness.")
add_label_paragraph(doc, "Target time  ", "3 minutes 15 seconds to 3 minutes 45 seconds")

add_heading(doc, "Speaking script", 2)
add_script(doc, [
    "The experiment uses four matched 100 millilitre thermal cores. Each starts near 80 degrees Celsius, while a separate probe records the actual ice-bath temperature. We record one reading every minute from zero to fifteen minutes. Consistent water volume, probe depth and immersion position make the comparison meaningful.",
    "We test a bare control, bubble wrap, Mylar alone and a full multilayer package made from cotton, bubble wrap and Mylar. The bare core shows the reference heat loss. Bubble wrap traps air. Mylar mainly reduces infrared exchange. The full package combines barriers against conduction, convection and radiation.",
    "We use Newton's cooling model to create a reference curve. The effective coefficient k describes the complete test system, including the core, insulation, geometry and boundary conditions. A smaller k means the core approaches the bath temperature more slowly. It is not the material's thermal conductivity.",
    "The method has four stages: prepare and check the equipment, condition every core to the agreed starting range, measure at fixed intervals and analyse the curves. Safety is part of the design. An adult handles water near 80 degrees Celsius, the bath stays stable and below waist height, electronics remain dry, and the test stops for leakage, probe movement or wet insulation.",
])

add_heading(doc, "Points to emphasize", 2)
add_bullet(doc, "Do not compress the cotton or bubble wrap because trapped air is part of the insulation mechanism.")
add_bullet(doc, "If starting temperatures differ, compare temperature drop, normalized retention or fitted k rather than final temperature alone.")
add_bullet(doc, "Safety and experimental control are engineering requirements, not side notes.")

add_heading(doc, "Handoff to Piyath", 2)
p = doc.add_paragraph("Once the physical test is controlled, the next step is to compare it with a prediction. Piyath will show how Mission Control connects the model with the students' measurements.")
p.paragraph_format.keep_together = True

doc.add_page_break()

# Presenter 3
add_heading(doc, "Presenter 3 Piyath", 1)
add_label_paragraph(doc, "Slides  ", "9 to 11")
add_label_paragraph(doc, "Goal  ", "Show the digital workflow clearly and complete the live demonstration without getting lost in interface details.")
add_label_paragraph(doc, "Target time  ", "3 minutes 45 seconds to 4 minutes 45 seconds")

add_heading(doc, "Speaking script", 2)
add_script(doc, [
    "Mission Control closes the loop between theory and evidence. Students choose an insulation configuration, enter the measured starting and bath temperatures, generate a predicted cooling curve and then enter the readings from the physical experiment.",
    "The thermometer remains manual by design. Students must observe and record every measurement, while the browser handles graphing and error calculations. This keeps them involved in the experiment and gives immediate feedback.",
    "For the demonstration, I will select the full multilayer configuration and confirm k equals 0.015 per minute. With an 80 degree starting temperature, a zero degree boundary and a fifteen-minute duration, the model predicts 63.9 degrees at the end.",
    "Next, I will open Measurements and show the supplied 16-point full-MLI series. Mission Control overlays the observations and reports a mean absolute error of 0.98 degrees Celsius. The residuals do not stay on one side of the curve, which reminds us that real measurements do not follow one perfect exponential.",
    "Under the common 80 degree and zero degree assumptions, the reference model separates the expected cooling rates. Full MLI retains the most heat, bubble wrap is next, and the bare and Mylar-only configurations cool more quickly in this classroom boundary.",
])

add_heading(doc, "Demo sequence", 2)
add_number(doc, "1.  Open Simulator and select Full MLI.")
add_number(doc, "2.  Confirm k = 0.015 min⁻¹, T₀ = 80 °C, Tenv = 0 °C and 15 minutes.")
add_number(doc, "3.  Run the model and point to T(15) = 63.9 °C.")
add_number(doc, "4.  Open Measurements, show 16 of 16 points and point to MAE = 0.98 °C.")

add_heading(doc, "Fallback and handoff", 2)
add_label_paragraph(doc, "If the live demo fails  ", "Stay on slide 10 and use the verified screenshots. Explain the same four steps without troubleshooting in front of the audience.")
p = doc.add_paragraph("The model gives us an expected ranking. Nuhansee will now compare it with the measured example and explain what the differences mean.")
p.paragraph_format.keep_together = True

doc.add_page_break()

# Presenter 4
add_heading(doc, "Presenter 4 Nuhansee", 1)
add_label_paragraph(doc, "Slides  ", "12 to 14")
add_label_paragraph(doc, "Goal  ", "State the measured result, interpret the unusual Mylar result correctly and finish with a clear explanation of what the project offers.")
add_label_paragraph(doc, "Target time  ", "2 minutes 45 seconds to 3 minutes 15 seconds")

add_heading(doc, "Speaking script", 2)
add_script(doc, [
    "In the worked example, the full multilayer package had the smallest temperature drop at 14.9 degrees Celsius and retained 74.1 percent of its heat above the ambient reference. Bubble wrap ranked second. The bare and Mylar-only samples cooled much more quickly.",
    "Because the trials did not start at exactly the same temperature, we should not rank them only by the final reading. Temperature drop and normalized heat retention give a fairer comparison. These values come from one worked example, so repeated measurements would be needed for validation.",
    "The Mylar-only result is especially useful. Reflective foil targets infrared radiation, which matters strongly in vacuum. In the classroom water bath, direct liquid contact, leakage and construction quality can dominate. The weak run therefore shows the importance of boundary conditions. It does not prove that reflective insulation is ineffective in space.",
    "SPHERE offers a complete learning workflow. Students predict, build, measure and critique the same thermal system. Teachers receive a reusable experiment, Mission Control, manuals, an illustrated procedure, a task sheet, flashcards and an aligned assessment. The result is a classroom-ready package that connects thermal theory with evidence and engineering judgment.",
    "Thank you. We welcome questions about the experiment, the model, Mission Control or classroom implementation.",
])

add_heading(doc, "Points to emphasize", 2)
add_bullet(doc, "Full MLI performed best in the supplied example, but the experiment still requires repeated trials for validation.")
add_bullet(doc, "Probe position, neck sealing, stratification, bath warming, immersion depth and layer compression can change the result.")
add_bullet(doc, "The closing offer is the integrated classroom package, not a claim that the kit reproduces a real spacesuit.")

add_heading(doc, "Contribution wording on slide 14", 2)
p = doc.add_paragraph("Read the contribution attribution exactly as it appears on the slide. Do not add or infer responsibilities that the project record does not document.")
p.paragraph_format.keep_together = True

doc.add_page_break()

# Team practice and Q&A
add_heading(doc, "Team handoffs and questions", 1)
add_heading(doc, "Handoff rules", 2)
add_bullet(doc, "The next presenter should already be standing and ready before the handoff sentence ends.")
add_bullet(doc, "The current presenter finishes the thought, pauses and says the next presenter's name.")
add_bullet(doc, "The next presenter begins with the connection to the previous section, not a new greeting.")
add_bullet(doc, "Only one person controls the computer during the live demo. Piyath owns the demo and fallback decision.")

add_heading(doc, "Question ownership", 2)
add_qa_table(doc)

add_heading(doc, "Useful short answers", 2)
add_label_paragraph(doc, "Is this a vacuum simulation  ", "No. It is a classroom analogy for thermal-control principles. The water bath emphasizes conduction and convection, while real spacecraft also depend strongly on radiation control and active systems.")
add_label_paragraph(doc, "Why did Mylar perform poorly  ", "The classroom boundary and build quality can let liquid conduction or leakage dominate. One weak run does not describe Mylar's performance in space.")
add_label_paragraph(doc, "Why enter data manually  ", "Manual entry keeps students responsible for observation while the web tool provides immediate graphing and model comparison.")
add_label_paragraph(doc, "What is the main value  ", "SPHERE combines theory, physical testing, software and assessment in one consistent classroom workflow.")

add_heading(doc, "Final rehearsal checklist", 2)
for item in [
    "Run the complete presentation once with a timer and keep the main section under 15 minutes.",
    "Practise all three handoff sentences exactly once, then make them sound natural.",
    "Verify the demo values: full MLI, k = 0.015 min⁻¹, T(15) = 63.9 °C and MAE = 0.98 °C.",
    "Open the fallback screenshots before presenting and confirm they are readable offline.",
    "Agree that the speaker who receives a question answers first; another person adds only if needed.",
    "Keep appendix slides 15 to 19 ready for apparatus, procedure, raw data, calculations and assessment questions.",
]:
    add_bullet(doc, "☐ " + item)

footer = doc.sections[0].footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("SPACESUIT THERMAL SHIELDING KIT  |  Four Presenter Guide")
fr.font.name = "Aptos"
fr.font.size = Pt(8.5)
fr.font.color.rgb = RGBColor(95, 104, 112)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
